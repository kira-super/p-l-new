from openpyxl import load_workbook
wb = load_workbook(r"C:\Users\kgontar\v2\oefof_pl\out\OAKS_EM_Stock_PL_Report.xlsx", read_only=True)
for sheet in ("AS", "VS", "JB", "KX", "SB", "HK", "IS"):
    ws = wb[sheet]
    swaps = []
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 9: continue
        if row[8] in ("SWAP", "FTSWAP"):
            swaps.append((row[1], row[8], row[2], row[4], row[5], row[6]))
    if swaps:
        print(f"\n=== {sheet} ({len(swaps)} swaps) ===")
        print(f"  {'Stock':<32} {'Type':<7} {'Total':>14} {'Realised':>14} {'Unrealised':>14} {'Income':>14}")
        for s in swaps[:8]:
            sname = (s[0] or "")[:32]
            print(f"  {sname:<32} {s[1]:<7} {s[2] or 0:>14.2f} {s[3] or 0:>14.2f} {s[4] or 0:>14.2f} {s[5] or 0:>14.2f}")
print(f"\nGross Exp KPI cols 9-10 widths in AS: {wb['AS'].column_dimensions['I'].width}, {wb['AS'].column_dimensions['J'].width}")
