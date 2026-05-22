from openpyxl import load_workbook
wb = load_workbook(r"C:\Users\kgontar\v2\oefof_pl\out\OAKS_EM_Stock_PL_Report.xlsx", read_only=True)
ws = wb["ValuationA"]
print(f"ValuationA: max_row={ws.max_row}, max_col={ws.max_column}")
for r_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if 1 <= r_i <= 10 or r_i in (12, 20, 50):
        print(f"  row {r_i}: {row}")

# Check swap rows in an analyst sheet
ws2 = wb["AS"]
print(f"\n--- AS sheet (max_row={ws2.max_row}) ---")
hdr = None
for r_i, row in enumerate(ws2.iter_rows(values_only=True), start=1):
    if row and "ISIN" in (row or ()):
        hdr = list(row); break
if hdr:
    # find Realised, Unrealised, Income, Total columns
    cols = {h: i for i, h in enumerate(hdr) if h}
    print("Cols:", {k: v for k, v in cols.items() if "P&L" in k or k == "Stock Name" or "Income" in k or "Instrument" in k})
    cnt = 0
    for r_i, row in enumerate(ws2.iter_rows(values_only=True), start=1):
        if row and len(row) > 5 and row[5] in ("SWAP", "FTSWAP"):
            sname = row[1]
            inst = row[5]
            r_pl = row[cols.get("Realised P&L (EUR)", -1)] if "Realised P&L (EUR)" in cols else "?"
            u_pl = row[cols.get("Unrealised P&L (EUR)", -1)] if "Unrealised P&L (EUR)" in cols else "?"
            i_pl = row[cols.get("Income (EUR)", -1)] if "Income (EUR)" in cols else "?"
            t_pl = row[cols.get("Total P&L (EUR)", -1)] if "Total P&L (EUR)" in cols else "?"
            print(f"  {sname[:32]:32s} {inst:6s} R={r_pl} U={u_pl} I={i_pl} T={t_pl}")
            cnt += 1
            if cnt > 12: break
