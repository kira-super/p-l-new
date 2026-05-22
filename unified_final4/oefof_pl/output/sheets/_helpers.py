"""Shared helper functions and constants for all sheet writers."""
from __future__ import annotations

from typing import Iterable, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ...compute.exposure import ExposureBlock, ExposureMetrics, PeriodMeta
from ..styles import (
    ALIGN_CENTER,
    ALIGN_KPI_LABEL,
    ALIGN_KPI_VALUE,
    ALIGN_LEFT,
    ALIGN_RIGHT,
    COLOR_BAND_ALT,
    COLOR_HEADER,
    COLOR_HEADER_ACCENT,
    COLOR_KPI_FILL,
    COLOR_KPI_FILL_ACCENT,
    COLOR_MUTED,
    COLOR_NEG,
    COLOR_POS,
    COLOR_TAB_ANALYST,
    COLOR_TAB_AUDIT,
    COLOR_TAB_DASHBOARD,
    COLOR_TAB_DATA,
    COLOR_VALIDATION_FAIL,
    COLOR_VALIDATION_PASS,
    COLOR_VALIDATION_WARN,
    COLUMN_LAYOUT_ANALYST,
    COLUMN_LAYOUT_HOLDINGS,
    COLUMN_LAYOUT_INCOME,
    ColSpec,
    FILL_BAND,
    FILL_HEADER,
    FILL_HEADER_ACCENT,
    FILL_KPI,
    FILL_KPI_ACCENT,
    FILL_SUBTOTAL,
    FILL_TOTAL,
    FONT_BODY,
    FONT_BODY_MUTED,
    FONT_HEADER,
    FONT_NEG,
    FONT_NEG_BOLD,
    FONT_POS,
    FONT_KPI_LABEL,
    FONT_KPI_VALUE,
    FONT_KPI_VALUE_NEG,
    FONT_KPI_VALUE_POS,
    FONT_SECTION,
    FONT_SUBTITLE,
    FONT_TITLE,
    FONT_TOTAL,
    align_for,
    fmt_for,
)

# Display names for known analyst codes.
ANALYST_DISPLAY: dict[str, str] = {
    "IS": "Ian Simmons",
    "HK": "Hayden Kwan",
    "JB": "Julius Bottcher",
    "SB": "Stefan Bottcher",
    "KX": "Karen Xiao",
    "VS": "Vijay Singh",
    "AS": "Alexander Short",
    "VJ": "VJ",
    "UNASSIGNED": "Unassigned",
}

# Stable sheet order for analysts.
ANALYST_TAB_ORDER: tuple[str, ...] = (
    "IS", "HK", "JB", "SB", "KX", "VS", "AS", "VJ", "UNASSIGNED",
)


# ─── Border constants ────────────────────────────────────────────────────────

_THIN = Side(style="thin", color="D5DDDE")
_MED = Side(style="medium", color=COLOR_HEADER)
BORDER_THIN_ALL = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
BORDER_TOP_THICK = Border(top=_MED)
BORDER_KPI = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# ─── Layout validation ───────────────────────────────────────────────────────

def _validate_layout_against_pl_df(pl_df: pd.DataFrame, layout: tuple[ColSpec, ...]) -> None:
    if pl_df.empty:
        return
    missing = [c.df_column for c in layout if c.df_column not in pl_df.columns]
    if missing:
        raise KeyError(
            f"Workbook layout references columns not in pl_df: {missing}. "
            "Either add them to pl_df or remove the ColSpec."
        )


# ─── Aggregation helpers ─────────────────────────────────────────────────────

def _current_holdings(pl_df: pd.DataFrame) -> pd.DataFrame:
    if pl_df.empty:
        return pl_df
    return pl_df[pl_df["Ending Units"] != 0].reset_index(drop=True)


def _exited_positions(pl_df: pd.DataFrame) -> pd.DataFrame:
    if pl_df.empty:
        return pl_df
    return pl_df[pl_df["Ending Units"] == 0].reset_index(drop=True)


def _reduced_positions(pl_df: pd.DataFrame) -> pd.DataFrame:
    """Positions still held at period end but with sells inside the period."""
    if pl_df.empty:
        return pl_df
    sold = pd.to_numeric(pl_df.get("Units Sold", 0.0), errors="coerce").fillna(0.0)
    held = pl_df["Ending Units"] != 0
    return pl_df[held & (sold > 0)].reset_index(drop=True)


def _with_total_contribution_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Add contribution-of-total and absolute-bps columns for display tables."""
    out = df.copy()
    if "Total P&L (EUR)" not in out.columns:
        out["Contribution to Total (%)"] = None
        out["Absolute Contribution (bps)"] = None
        return out

    total_pl = float(pd.to_numeric(out["Total P&L (EUR)"], errors="coerce").fillna(0.0).sum())
    if abs(total_pl) < 1e-9:
        contrib = pd.Series([None] * len(out), index=out.index, dtype=object)
    else:
        contrib = pd.to_numeric(out["Total P&L (EUR)"], errors="coerce") / total_pl
    out["Contribution to Total (%)"] = contrib
    out["Absolute Contribution (bps)"] = pd.to_numeric(contrib, errors="coerce").abs() * 10_000
    return out


def _analyst_order(pl_df: pd.DataFrame) -> list[str]:
    """Distinct analyst codes in stable display order."""
    if pl_df.empty or "Analyst" not in pl_df.columns:
        return []
    found = {str(a).strip().upper() for a in pl_df["Analyst"].dropna() if str(a).strip()}
    ordered = [a for a in ANALYST_TAB_ORDER if a in found]
    extras = sorted(found - set(ordered))
    return ordered + extras


def _analyst_display(code: str) -> str:
    return ANALYST_DISPLAY.get(code, code)


def _safe_float(v) -> float:
    try:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return 0.0
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _pct_or_none(numerator: float, denominator: float) -> float | None:
    """Ratio (NOT percent). Returns None if denominator is ~0 to avoid div-by-zero."""
    if abs(denominator) < 1e-9:
        return None
    return numerator / denominator


def _exposure_for(df: pd.DataFrame) -> dict:
    """Compute long/short/net exposure (EUR) and P&L breakdown."""
    if df.empty:
        return dict(
            n=0, long_eur=0.0, short_eur=0.0, net_eur=0.0, gross_eur=0.0,
            realised=0.0, unrealised=0.0, income=0.0, fx_attrib=0.0,
            total=0.0, total_pct=None, cost=0.0,
        )
    mv = pd.to_numeric(df.get("Market Value End (EUR)", 0.0), errors="coerce").fillna(0.0)
    cost = pd.to_numeric(df.get("Cost Basis (EUR)", 0.0), errors="coerce").fillna(0.0).abs()
    instr = df.get("Instrument", "ORD").astype(str).str.upper()
    ls = df.get("L/S", "L").astype(str).str.upper()
    is_deriv = instr.isin(["SWAP", "FTSWAP"])
    sign = ls.map(lambda s: -1.0 if s == "S" else 1.0)
    exposure_signed = (cost.where(is_deriv, mv.abs())) * sign
    long_eur = float(exposure_signed[exposure_signed > 0].sum())
    short_eur = abs(float(exposure_signed[exposure_signed < 0].sum()))
    realised = float(pd.to_numeric(df.get("Realised P&L (EUR)", 0.0), errors="coerce").fillna(0.0).sum())
    unrealised = float(pd.to_numeric(df.get("Unrealised P&L (EUR)", 0.0), errors="coerce").fillna(0.0).sum())
    income = float(pd.to_numeric(df.get("Income (EUR)", 0.0), errors="coerce").fillna(0.0).sum())
    total = float(pd.to_numeric(df.get("Total P&L (EUR)", 0.0), errors="coerce").fillna(0.0).sum())
    fx_attrib = total - (realised + unrealised + income)
    cost_total = float(cost.sum())
    return dict(
        n=int(len(df)),
        long_eur=long_eur, short_eur=short_eur,
        net_eur=long_eur - short_eur,
        gross_eur=long_eur + short_eur,
        realised=realised, unrealised=unrealised, income=income,
        fx_attrib=fx_attrib,
        total=total, total_pct=_pct_or_none(total, cost_total),
        cost=cost_total,
    )


# ─── Common sheet helpers ────────────────────────────────────────────────────

def _setup_sheet(ws: Worksheet, *, tab_color: str | None = None,
                 freeze: str | None = None, hide_gridlines: bool = True,
                 zoom: int = 100) -> None:
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    if hide_gridlines:
        ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    if freeze:
        ws.freeze_panes = freeze


def _write_title_block(ws: Worksheet, meta: PeriodMeta, *,
                       title: str | None = None, subtitle: str | None = None,
                       col_start: int = 1) -> None:
    title = title or f"{meta.fund_name} — Stock-Level P&L Report"
    subtitle = subtitle or (
        f"Period: {meta.start_date.date():%d %B %Y} → "
        f"{meta.end_date.date():%d %B %Y}  |  Fund Currency: EUR"
    )
    ws.row_dimensions[1].height = 30
    c = ws.cell(1, col_start, title)
    c.font = FONT_TITLE
    c.alignment = ALIGN_LEFT
    c2 = ws.cell(2, col_start, subtitle)
    c2.font = FONT_BODY_MUTED
    c2.alignment = ALIGN_LEFT


def _write_section_header(ws: Worksheet, row: int, *, title: str,
                          span_cols: int, accent: bool = False,
                          col_start: int = 1) -> None:
    fill = FILL_HEADER_ACCENT if accent else FILL_HEADER
    ws.row_dimensions[row].height = 22
    for c in range(col_start, col_start + span_cols):
        cell = ws.cell(row, c)
        cell.fill = fill
        cell.font = FONT_SECTION
        cell.alignment = ALIGN_LEFT
    ws.cell(row, col_start, "  " + title).font = FONT_SECTION


def _kpi_card(ws: Worksheet, *, row: int, col: int, label: str, value,
              kind: str = "eur", colour_signed: bool = False,
              accent: bool = False) -> None:
    """Render a 2-row × 2-col card: label on top, value below."""
    fill = FILL_KPI_ACCENT if accent else FILL_KPI

    ws.merge_cells(start_row=row, start_column=col,
                   end_row=row, end_column=col + 1)
    lc = ws.cell(row, col, label.upper())
    lc.font = FONT_KPI_LABEL
    lc.alignment = ALIGN_KPI_LABEL
    lc.fill = fill
    lc.border = BORDER_KPI
    ws.cell(row, col + 1).fill = fill
    ws.cell(row, col + 1).border = BORDER_KPI

    ws.merge_cells(start_row=row + 1, start_column=col,
                   end_row=row + 1, end_column=col + 1)
    vc = ws.cell(row + 1, col, _coerce_value(value, kind))
    vc.alignment = ALIGN_KPI_VALUE
    vc.fill = fill
    vc.border = BORDER_KPI
    if value is None:
        vc.font = FONT_KPI_VALUE
    elif colour_signed and isinstance(vc.value, (int, float)):
        vc.font = FONT_KPI_VALUE_POS if vc.value >= 0 else FONT_KPI_VALUE_NEG
    else:
        vc.font = FONT_KPI_VALUE
    if kind != "text":
        vc.number_format = fmt_for(kind)
    ws.cell(row + 1, col + 1).fill = fill
    ws.cell(row + 1, col + 1).border = BORDER_KPI

    ws.row_dimensions[row].height = 18
    ws.row_dimensions[row + 1].height = 26


def _write_exposure_block_table(
    ws: Worksheet,
    block: ExposureBlock,
    analysts: list[str],
    meta: PeriodMeta,
    *,
    start_row: int,
    is_ytd: bool = False,
    col_start: int = 1,
    show_total: bool = True,
) -> int:
    """Write a current-snapshot or time-weighted average exposure table.

    Returns the next available row after the table.
    """
    if block.as_of is not None:
        date_label = "through" if is_ytd else "as of"
        as_of_str = f"  ({date_label} {block.as_of.date():%d %B %Y})"
    else:
        as_of_str = ""
    title = block.label + as_of_str
    _write_section_header(ws, start_row, title=title, span_cols=9,
                          accent=is_ytd, col_start=col_start)
    start_row += 1

    prefix = "YTD Avg " if is_ytd else "Current "
    hdrs = (
        "Analyst",
        f"{prefix}Long ({meta.fund_currency})", f"{prefix}Long (%)",
        f"{prefix}Short ({meta.fund_currency})", f"{prefix}Short (%)",
        f"{prefix}Net ({meta.fund_currency})", f"{prefix}Net (%)",
        f"{prefix}Gross ({meta.fund_currency})", f"{prefix}Gross (%)",
    )
    kinds = ("text", "eur", "pct", "eur", "pct", "eur", "pct", "eur", "pct")
    hdr_fill = FILL_HEADER_ACCENT if is_ytd else FILL_HEADER
    for i, h in enumerate(hdrs, start=col_start):
        cc = ws.cell(start_row, i, h)
        cc.font = FONT_HEADER; cc.fill = hdr_fill; cc.alignment = ALIGN_CENTER
    start_row += 1

    for r_idx, code in enumerate(analysts):
        em = block.by_analyst.get(code, ExposureMetrics())
        vals = (
            _analyst_display(code),
            em.long_eur, em.pct_of_nav(em.long_eur),
            em.short_eur, em.pct_of_nav(em.short_eur),
            em.net_eur, em.pct_of_nav(em.net_eur),
            em.gross_eur, em.pct_of_nav(em.gross_eur),
        )
        for c_idx, (v, k) in enumerate(zip(vals, kinds), start=col_start):
            cell = ws.cell(start_row + r_idx, c_idx, _coerce_value(v, k))
            cell.font = FONT_BODY
            if r_idx % 2 == 1:
                cell.fill = FILL_BAND
            cell.alignment = align_for(k)
            if k != "text":
                cell.number_format = fmt_for(k)

    last_data_row = start_row + len(analysts) - 1

    if show_total:
        em = block.fund
        total_row = last_data_row + 1
        vals = (
            "Fund Total",
            em.long_eur, em.pct_of_nav(em.long_eur),
            em.short_eur, em.pct_of_nav(em.short_eur),
            em.net_eur, em.pct_of_nav(em.net_eur),
            em.gross_eur, em.pct_of_nav(em.gross_eur),
        )
        for c_idx, (v, k) in enumerate(zip(vals, kinds), start=col_start):
            cell = ws.cell(total_row, c_idx, _coerce_value(v, k))
            cell.font = FONT_TOTAL; cell.fill = FILL_TOTAL
            cell.alignment = align_for(k)
            if k != "text":
                cell.number_format = fmt_for(k)
        return total_row + 1

    return last_data_row + 1


# ─── Generic table writers ───────────────────────────────────────────────────

def _write_layout_headers(ws: Worksheet, row: int, layout: tuple[ColSpec, ...],
                          *, accent: bool = False,
                          suppress: frozenset[str] = frozenset()) -> None:
    fill = FILL_HEADER_ACCENT if accent else FILL_HEADER
    ws.row_dimensions[row].height = 24
    for col_idx, spec in enumerate(layout, start=1):
        label = "" if spec.df_column in suppress else (spec.display or spec.df_column)
        cell = ws.cell(row, col_idx, label)
        cell.font = FONT_HEADER
        cell.fill = fill
        cell.alignment = ALIGN_CENTER


def _write_layout_data(ws: Worksheet, df: pd.DataFrame,
                       layout: tuple[ColSpec, ...], *, start_row: int,
                       suppress: frozenset[str] = frozenset()) -> int:
    for row_offset, (_, row) in enumerate(df.iterrows()):
        excel_row = start_row + row_offset
        banded = (row_offset % 2 == 1)
        for col_idx, spec in enumerate(layout, start=1):
            if spec.df_column in suppress:
                cell = ws.cell(excel_row, col_idx)
                cell.font = FONT_BODY
                if banded:
                    cell.fill = FILL_BAND
                continue
            value = row.get(spec.df_column, None) if spec.df_column in df.columns else None
            cell = ws.cell(excel_row, col_idx, _coerce_value(value, spec.kind))
            cell.font = FONT_BODY
            cell.number_format = fmt_for(spec.kind)
            cell.alignment = align_for(spec.kind)
            if banded:
                cell.fill = FILL_BAND
            if spec.kind == "eur" and isinstance(cell.value, (int, float)):
                if "P&L" in spec.df_column or "Income" in spec.df_column:
                    if cell.value < 0:
                        cell.font = FONT_NEG
                    elif cell.value > 0 and "P&L" in spec.df_column:
                        cell.font = FONT_POS
    return start_row + len(df)


def _write_layout_total(ws: Worksheet, df: pd.DataFrame,
                        layout: tuple[ColSpec, ...], *, total_row: int,
                        label: str = "Total", fill: PatternFill = FILL_TOTAL,
                        suppress: frozenset[str] = frozenset()) -> int:
    ws.row_dimensions[total_row].height = 20
    for col_idx, spec in enumerate(layout, start=1):
        cell = ws.cell(total_row, col_idx)
        cell.font = FONT_TOTAL
        cell.fill = fill
        cell.alignment = align_for(spec.kind)
        cell.border = BORDER_TOP_THICK
        if spec.df_column in suppress:
            continue
        if spec.total == "label" or col_idx == 1:
            cell.value = label
        elif spec.total == "sum" and spec.df_column in df.columns:
            s = pd.to_numeric(df[spec.df_column], errors="coerce").sum()
            cell.value = float(s)
            cell.number_format = fmt_for(spec.kind)
            if spec.kind == "eur" and float(s) < 0:
                cell.font = FONT_NEG_BOLD
        elif spec.df_column == "Contribution to Total (%)":
            pl = pd.to_numeric(df.get("Total P&L (EUR)",
                                      pd.Series(dtype=float)),
                               errors="coerce").sum()
            if abs(float(pl)) > 1e-9:
                cell.value = 1.0
                cell.number_format = fmt_for(spec.kind)
        elif spec.df_column == "Total P&L (%)":
            pl = pd.to_numeric(df.get("Total P&L (EUR)",
                                       pd.Series(dtype=float)),
                               errors="coerce").sum()
            cb = pd.to_numeric(df.get("Cost Basis (EUR)",
                                       pd.Series(dtype=float)),
                               errors="coerce").sum()
            if cb:
                ratio = float(pl) / abs(float(cb))
                cell.value = ratio
                cell.number_format = fmt_for(spec.kind)
                if ratio < 0:
                    cell.font = FONT_NEG_BOLD
    return total_row + 1


def _write_table(wb: Workbook, sheet_name: str, df: pd.DataFrame,
                 layout: tuple[ColSpec, ...], meta: PeriodMeta,
                 *, tab_color: str | None = None) -> None:
    ws = wb.create_sheet(sheet_name)
    _setup_sheet(ws, tab_color=tab_color)
    _write_title_block(ws, meta,
                       title=f"{sheet_name} — {meta.fund_name}",
                       subtitle=(f"Period: {meta.start_date.date():%d %B %Y} → "
                                 f"{meta.end_date.date():%d %B %Y}"))
    n_cols = len(layout)
    for i, spec in enumerate(layout, start=1):
        ws.column_dimensions[get_column_letter(i)].width = spec.width

    if df.empty:
        ws.cell(4, 1, "(no rows)").font = FONT_BODY_MUTED
        return

    if any(
        spec.df_column in {"Contribution to Total (%)", "Absolute Contribution (bps)"}
        for spec in layout
    ):
        df = _with_total_contribution_columns(df)

    if sheet_name == "Current Holdings":
        _validate_layout_against_pl_df(df, COLUMN_LAYOUT_HOLDINGS)

    if "Total P&L (EUR)" in df.columns:
        df = df.sort_values("Total P&L (EUR)", ascending=False, na_position="last").reset_index(drop=True)

    header_row = 4
    _write_layout_headers(ws, header_row, layout)

    if sheet_name in ("Exited Positions", "Reduced Positions"):
        total_header_row = header_row + 1
        ws.merge_cells(start_row=total_header_row, start_column=2,
                       end_row=total_header_row, end_column=n_cols)
        total_cell = ws.cell(total_header_row, 2, "Total")
        total_cell.font = FONT_TOTAL
        total_cell.fill = FILL_TOTAL
        total_cell.alignment = ALIGN_CENTER
        ws.row_dimensions[total_header_row].height = 20
        next_row = total_header_row + 1
    else:
        next_row = header_row + 1

    next_row = _write_layout_data(ws, df, layout, start_row=next_row)
    if any(spec.total != "blank" for spec in layout):
        next_row = _write_layout_total(ws, df, layout, total_row=next_row,
                                        label="Total", fill=FILL_TOTAL)

    ws.freeze_panes = ws.cell(header_row + 1, 1)

    if sheet_name in ("Exited Positions", "Reduced Positions"):
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{header_row + 1 + len(df)}"
    else:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{header_row + len(df)}"


# ─── Coercion / sheet naming ─────────────────────────────────────────────────

def _coerce_value(v, kind: str):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    if kind == "date":
        if isinstance(v, pd.Timestamp):
            if pd.isna(v):
                return None
            return v.to_pydatetime()
        return v
    if kind in ("int", "float", "price", "eur", "pct"):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    return v if isinstance(v, str) else str(v)


def _safe_sheet_name(name: str) -> str:
    """Excel sheet names: max 31 chars; cannot contain : \\ / ? * [ ]."""
    bad = ':\\/?*[]'
    cleaned = "".join("_" if ch in bad else ch for ch in str(name))
    return cleaned[:31] or "Sheet"
