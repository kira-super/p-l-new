"""Summary dashboard sheet writer."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from ...compute.exposure import ExposureMetrics, PeriodMeta
from ..styles import (
    ALIGN_CENTER,
    ALIGN_LEFT,
    COLOR_TAB_DASHBOARD,
    FILL_BAND,
    FILL_HEADER,
    FILL_HEADER_ACCENT,
    FILL_TOTAL,
    FONT_BODY,
    FONT_BODY_MUTED,
    FONT_HEADER,
    FONT_NEG,
    FONT_NEG_BOLD,
    FONT_TOTAL,
    align_for,
    fmt_for,
)
from ._helpers import (
    _analyst_display,
    _analyst_order,
    _normalise_analyst_code,
    _coerce_value,
    _current_holdings,
    _exited_positions,
    _exposure_for,
    _kpi_card,
    _pct_or_none,
    _setup_sheet,
    _write_exposure_block_table,
    _write_section_header,
    _write_title_block,
    get_column_letter,
)


def _write_summary(wb: Workbook, pl_df: pd.DataFrame, income_df: pd.DataFrame,
                   meta: PeriodMeta) -> None:
    ws = wb.create_sheet("Summary")
    _setup_sheet(ws, tab_color=COLOR_TAB_DASHBOARD, zoom=100)
    _write_title_block(ws, meta)

    summary_widths = [18, 11, 17, 13, 21, 22, 20, 22, 14, 16, 26, 22]
    for i, w in enumerate(summary_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    if pl_df.empty:
        ws.cell(4, 1, "No positions in this period.").font = FONT_BODY_MUTED
        return

    current = _current_holdings(pl_df)
    exited = _exited_positions(pl_df)
    fund_kpi = _exposure_for(pl_df)
    cur_kpi = _exposure_for(current)
    ex_kpi = _exposure_for(exited)
    ytd_exp = meta.ytd_weighted_exposure

    card_rows = (4, 6, 8)
    card_cols = (1, 3, 5)

    cards = [
        ("Total Positions", len(pl_df), "int", False, False),
        ("Current Holdings", cur_kpi["n"], "int", False, False),
        ("Exited Positions", ex_kpi["n"], "int", False, False),

        ("Realised P&L (EUR)", fund_kpi["realised"], "eur", True, False),
        ("Unrealised P&L (EUR)", fund_kpi["unrealised"], "eur", True, False),
        ("Income & Financing (EUR)", fund_kpi["income"], "eur", True, False),

        ("FX Attribution (EUR)", fund_kpi["fx_attrib"], "eur", True, False),
        ("Total P&L (EUR)", fund_kpi["total"], "eur", True, True),
    ]
    for i, (label, value, kind, signed, accent) in enumerate(cards):
        r = card_rows[i // 3]
        c = card_cols[i % 3]
        _kpi_card(ws, row=r, col=c, label=label, value=value,
                  kind=kind, colour_signed=signed, accent=accent)

    base_row = card_rows[-1] + 3
    _write_section_header(ws, base_row,
                          title="By Analyst — Period P&L Attribution (YTD)",
                          span_cols=8, accent=True)
    base_row += 1

    ret_pct_label = ("Return % (YTD Gross Exp)" if ytd_exp is not None
                     else "Total P&L % (Cost Basis)")
    a_headers = ("Analyst", "Open Positions",
                 "Total P&L (EUR)", ret_pct_label,
                 "Realised (EUR)", "Unrealised (EUR)",
                 "Income & Financing (EUR)", "FX Attribution (EUR)")
    a_kinds = ("text", "int", "eur", "pct", "eur", "eur", "eur", "eur")
    for i, h in enumerate(a_headers, start=1):
        cc = ws.cell(base_row, i, h); cc.font = FONT_HEADER
        cc.fill = FILL_HEADER_ACCENT; cc.alignment = ALIGN_CENTER
    base_row += 1

    def _ret_pct(total: float, code: str) -> float | None:
        if ytd_exp is not None:
            em = ytd_exp.by_analyst.get(code, ExposureMetrics())
            return _pct_or_none(total, em.gross_eur) if abs(em.gross_eur) > 1e-9 else None
        return None

    a_rows_with_codes: list[tuple[str, tuple]] = []
    for code in _analyst_order(pl_df, meta.analyst_codes):
        analyst_codes = pl_df["Analyst"].map(_normalise_analyst_code)
        sub = pl_df[analyst_codes == code]
        cur_sub = sub[sub["Ending Units"] != 0]
        e = _exposure_for(sub)
        ret = _ret_pct(e["total"], code) if ytd_exp is not None else e["total_pct"]
        row = (_analyst_display(code, meta.analyst_codes), int(len(cur_sub)),
               e["total"], ret,
               e["realised"], e["unrealised"], e["income"], e["fx_attrib"])
        a_rows_with_codes.append((code, row))

    a_rows_with_codes.sort(key=lambda x: x[1][2] if isinstance(x[1][2], (int, float)) else 0,
                           reverse=True)

    sorted_analyst_codes = [code for code, _ in a_rows_with_codes]
    a_rows = [row for _, row in a_rows_with_codes]

    fund_ret = (_pct_or_none(fund_kpi["total"], ytd_exp.fund.gross_eur)
                if ytd_exp is not None and abs(ytd_exp.fund.gross_eur) > 1e-9
                else fund_kpi["total_pct"])
    a_rows.append(("TOTAL", int(len(_current_holdings(pl_df))),
                   fund_kpi["total"], fund_ret,
                   fund_kpi["realised"], fund_kpi["unrealised"],
                   fund_kpi["income"], fund_kpi["fx_attrib"]))
    for r_idx, row in enumerate(a_rows):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(base_row + r_idx, c_idx,
                           _coerce_value(val, a_kinds[c_idx - 1]))
            kind = a_kinds[c_idx - 1]
            if r_idx == len(a_rows) - 1:
                cell.font = FONT_TOTAL; cell.fill = FILL_TOTAL
            else:
                cell.font = FONT_BODY
                if r_idx % 2 == 1:
                    cell.fill = FILL_BAND
            cell.alignment = align_for(kind)
            if kind != "text":
                cell.number_format = fmt_for(kind)
            if kind in ("eur", "pct") and isinstance(cell.value, (int, float)) and cell.value < 0:
                cell.font = FONT_NEG_BOLD if cell.font is FONT_TOTAL else FONT_NEG

    end_row = base_row + len(a_rows)
    analysts = sorted_analyst_codes

    sec_row = end_row + 2
    _write_section_header(ws, sec_row,
                          title="Current Portfolio Snapshot",
                          span_cols=9, accent=False)
    sec_row += 1

    if meta.snapshot_exposure is not None:
        sec_row = _write_exposure_block_table(
            ws, meta.snapshot_exposure, analysts, meta, start_row=sec_row, is_ytd=False,
        ) + 1
    else:
        ws.cell(sec_row, 1,
                "Current snapshot unavailable — end-period snapshot not loaded."
                ).font = FONT_BODY_MUTED
        sec_row += 2
    if meta.ytd_weighted_exposure is not None:
        ytd_fund_gross = meta.ytd_weighted_exposure.fund.gross_eur
        sec_row = _write_exposure_block_table(
            ws,
            meta.ytd_weighted_exposure,
            analysts,
            meta,
            start_row=sec_row,
            is_ytd=True,
            pct_denominator=ytd_fund_gross if abs(ytd_fund_gross) > 1e-9 else None,
        ) + 1
        _write_section_header(ws, sec_row, title="Return on Gross Exposure (YTD)",
                              span_cols=5)
        sec_row += 1
        ren_hdrs = (
            "Analyst",
            "Total P&L (EUR)",
            "YTD Avg Gross Exp (EUR)",
            "Return %",
            "Contribution to Fund Return",
        )
        ren_kinds = ("text", "eur", "eur", "pct", "pct")
        for i, h in enumerate(ren_hdrs, start=1):
            cc = ws.cell(sec_row, i, h)
            cc.font = FONT_HEADER; cc.fill = FILL_HEADER; cc.alignment = ALIGN_CENTER
        sec_row += 1
        ytd = meta.ytd_weighted_exposure
        fund_gross_eur = ytd.fund.gross_eur
        fund_total_pl = fund_kpi["total"]

        ren_rows: list[tuple[str, float, float, float | None, float | None]] = []
        for code in analysts:
            analyst_codes = pl_df["Analyst"].map(_normalise_analyst_code)
            sub = pl_df[analyst_codes == code]
            pl_total = float(pd.to_numeric(sub.get("Total P&L (EUR)", 0.0),
                                           errors="coerce").fillna(0.0).sum())
            em = ytd.by_analyst.get(code, ExposureMetrics())
            ret = _pct_or_none(pl_total, em.gross_eur) if abs(em.gross_eur) > 1e-9 else None
            fund_weighted_ret = (
                _pct_or_none(pl_total, fund_total_pl) if abs(fund_total_pl) > 1e-9 else None
            )
            ren_rows.append((_analyst_display(code, meta.analyst_codes), pl_total, em.gross_eur, ret, fund_weighted_ret))

        ren_rows.sort(key=lambda x: (x[3] is None, -x[3] if x[3] is not None else 0))

        for r_idx, vals in enumerate(ren_rows):
            for c_idx, (v, k) in enumerate(zip(vals, ren_kinds), start=1):
                cell = ws.cell(sec_row + r_idx, c_idx, _coerce_value(v, k))
                cell.font = FONT_BODY
                if r_idx % 2 == 1:
                    cell.fill = FILL_BAND
                cell.alignment = align_for(k)
                if k != "text":
                    cell.number_format = fmt_for(k)
                if k in ("eur", "pct") and isinstance(cell.value, (int, float)) and cell.value < 0:
                    cell.font = FONT_NEG

        pl_fund = fund_kpi["total"]
        em_fund = ytd.fund
        ret_fund = None
        fund_weighted_ret_fund = (
            _pct_or_none(pl_fund, fund_total_pl) if abs(fund_total_pl) > 1e-9 else None
        )
        for c_idx, (v, k) in enumerate(zip(
            ("TOTAL", pl_fund, em_fund.gross_eur, ret_fund, fund_weighted_ret_fund), ren_kinds
        ), start=1):
            cell = ws.cell(sec_row + len(ren_rows), c_idx, _coerce_value(v, k))
            cell.font = FONT_TOTAL; cell.fill = FILL_TOTAL
            cell.alignment = align_for(k)
            if k != "text":
                cell.number_format = fmt_for(k)
            if k in ("eur", "pct") and isinstance(cell.value, (int, float)) and cell.value < 0:
                cell.font = FONT_NEG_BOLD
        sec_row += len(ren_rows) + 2
    elif meta.snapshot_exposure is not None:
        ws.cell(sec_row, 1,
            "YTD gross-exposure history unavailable — snapshot history not loaded."
                ).font = FONT_BODY_MUTED
        sec_row += 2

    foot_row = sec_row + 1
    fc = ws.cell(foot_row, 1,
                 f"Generated {meta.run_timestamp:%Y-%m-%d %H:%M:%S}  |  "
                 f"Snapshot: {Path(meta.snapshot_path).name}  |  "
                 f"Bottler: {Path(meta.bottler_path).name}")
    fc.font = FONT_BODY_MUTED
