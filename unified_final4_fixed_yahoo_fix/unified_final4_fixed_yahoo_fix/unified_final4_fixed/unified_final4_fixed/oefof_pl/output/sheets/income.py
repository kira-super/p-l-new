"""Income summary and trade detail sheet writers."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook

from ...compute.exposure import PeriodMeta
from ..styles import (
    COLOR_TAB_DATA,
    ColSpec,
    FONT_BODY_MUTED,
)
from ._helpers import (
    _setup_sheet,
    _write_layout_data,
    _write_layout_headers,
    _write_layout_total,
    _write_title_block,
    get_column_letter,
)


def _write_income_summary(wb: Workbook, pl_df: pd.DataFrame,
                          meta: PeriodMeta) -> None:
    ws = wb.create_sheet("Income & Dividends")
    _setup_sheet(ws, tab_color=COLOR_TAB_DATA)
    _write_title_block(ws, meta,
                       title=f"Income & Dividends — {meta.fund_name}",
                       subtitle="By position")

    if pl_df.empty:
        ws.cell(4, 1, "(no positions)").font = FONT_BODY_MUTED
        return

    sub = pl_df[pl_df["Income (EUR)"] != 0][[
        "ISIN", "Stock Name", "Analyst", "CCY",
        "Income (Local)", "Income (EUR)", "Income Yield (%)",
    ]].reset_index(drop=True)

    layout = (
        ColSpec("ISIN",              14, "text",  total="label"),
        ColSpec("Stock Name",        32, "text"),
        ColSpec("Analyst",           10, "text"),
        ColSpec("CCY",                6, "text"),
        ColSpec("Income (Local)",    14, "float"),
        ColSpec("Income (EUR)",      14, "eur",   total="sum",
                display="Income & Financing (EUR)"),
        ColSpec("Income Yield (%)",  14, "pct"),
    )
    n_cols = len(layout)
    for i, spec in enumerate(layout, start=1):
        ws.column_dimensions[get_column_letter(i)].width = spec.width

    if sub.empty:
        ws.cell(4, 1, "(no income)").font = FONT_BODY_MUTED
        return

    header_row = 4
    _write_layout_headers(ws, header_row, layout)
    next_row = header_row + 1
    next_row = _write_layout_data(ws, sub, layout, start_row=next_row)
    next_row = _write_layout_total(ws, sub, layout, total_row=next_row)
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{header_row + len(sub)}"


def _write_trade_detail(wb: Workbook, trades_df: pd.DataFrame,
                        meta: PeriodMeta) -> None:
    ws = wb.create_sheet("Trade Detail")
    _setup_sheet(ws, tab_color=COLOR_TAB_DATA)
    _write_title_block(ws, meta,
                       title=f"Trade Detail — {meta.fund_name}",
                       subtitle="All trades in the analysis window")

    if trades_df.empty:
        ws.cell(4, 1, "(no trades)").font = FONT_BODY_MUTED
        return

    cols = ["TRADE_ID", "PCODE_ORIG", "ISIN", "SNAME", "CCY", "T", "CDATE",
            "UNITS", "GROSSPRICE_LOCAL", "BUY_NET_LOCAL", "SELL_NET_LOCAL",
            "INCOME_LOCAL"]
    sub = trades_df[[c for c in cols if c in trades_df.columns]].copy()
    layout = tuple(
        ColSpec(c, 14, "date" if c == "CDATE"
                       else "int" if c == "TRADE_ID"
                       else "float" if c not in ("PCODE_ORIG", "ISIN", "SNAME",
                                                  "CCY", "T")
                       else "text",
                total="blank")
        for c in sub.columns
    )
    n_cols = len(layout)
    for i, spec in enumerate(layout, start=1):
        ws.column_dimensions[get_column_letter(i)].width = spec.width
    header_row = 4
    _write_layout_headers(ws, header_row, layout)
    next_row = header_row + 1
    next_row = _write_layout_data(ws, sub, layout, start_row=next_row)
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{header_row + len(sub)}"
