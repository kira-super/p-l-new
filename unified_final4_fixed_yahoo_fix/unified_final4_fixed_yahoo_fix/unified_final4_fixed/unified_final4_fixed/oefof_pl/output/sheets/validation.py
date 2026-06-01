"""Validation report and methodology sheet writers."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from ...compute.exposure import PeriodMeta
from ...validate.checks import ValidationResult
from ..styles import (
    ALIGN_CENTER,
    ALIGN_LEFT,
    COLOR_TAB_AUDIT,
    COLOR_VALIDATION_FAIL,
    COLOR_VALIDATION_PASS,
    COLOR_VALIDATION_WARN,
    FILL_HEADER,
    FONT_BODY,
    FONT_HEADER,
    FONT_SUBTITLE,
    FONT_TOTAL,
    fmt_for,
)
from ._helpers import (
    _setup_sheet,
    _write_title_block,
    get_column_letter,
)


def _write_validation(
    wb: Workbook, validation: ValidationResult, meta: PeriodMeta,
    write_seconds: float,
    *, ca_suggestions_path: Path | None = None,
) -> None:
    ws = wb.create_sheet("Validation Report")
    _setup_sheet(ws, tab_color=COLOR_TAB_AUDIT)
    _write_title_block(ws, meta,
                       title="Validation Report",
                       subtitle=f"Workbook write took {write_seconds:.2f}s")

    if validation.has_critical:
        head = "STATUS: CRITICAL FAILURES PRESENT"
        col = COLOR_VALIDATION_FAIL
    elif validation.has_warnings:
        head = "STATUS: PASS WITH WARNINGS"
        col = COLOR_VALIDATION_WARN
    else:
        head = "STATUS: ALL CHECKS PASS"
        col = COLOR_VALIDATION_PASS
    c = ws.cell(3, 1, head); c.font = Font(bold=True, size=12, color=col)

    headers = ("ID", "Check", "Status", "Detail", "Failures")
    header_row = 5
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(header_row, i, h)
        cell.font = FONT_HEADER; cell.fill = FILL_HEADER; cell.alignment = ALIGN_CENTER

    r = header_row + 1
    for chk in validation.checks:
        ws.cell(r, 1, chk.id).font = FONT_BODY
        ws.cell(r, 2, chk.name).font = FONT_BODY
        sc = ws.cell(r, 3, chk.status)
        sc.font = Font(
            bold=True,
            color={"PASS": COLOR_VALIDATION_PASS,
                   "WARN": COLOR_VALIDATION_WARN,
                   "FAIL": COLOR_VALIDATION_FAIL}[chk.status],
        )
        ws.cell(r, 4, chk.detail).font = FONT_BODY
        ws.cell(r, 5, len(chk.failures)).font = FONT_BODY
        r += 1

    r += 2
    for chk in validation.checks:
        if not chk.failures:
            continue
        ws.cell(r, 1, f"{chk.id} — {chk.name}: {len(chk.failures)} failure(s)").font = FONT_SUBTITLE
        r += 1
        keys = list(chk.failures[0].keys())
        for i, k in enumerate(keys, start=1):
            cc = ws.cell(r, i, k)
            cc.font = FONT_HEADER; cc.fill = FILL_HEADER; cc.alignment = ALIGN_CENTER
        r += 1
        for f in chk.failures:
            for i, k in enumerate(keys, start=1):
                v = f.get(k, "")
                cell = ws.cell(r, i, v); cell.font = FONT_BODY
                if isinstance(v, float):
                    cell.number_format = fmt_for("float")
            r += 1
        r += 2

    try:
        val01 = validation.by_id("VAL-01")
    except KeyError:
        val01 = None
    if val01 is not None and val01.status == "FAIL" and val01.failures:
        cell = ws.cell(r, 1, "VAL-01 — How to unblock")
        cell.font = FONT_SUBTITLE
        r += 1
        msg_lines = [
            "Each VAL-01 failure means HiPort end-units do NOT match start + bought - sold.",
            "Causes are almost always corporate actions (bonus, split, consolidation, in-kind delivery)",
            "that hit the snapshot but not the Bottler trade feed.",
            "",
            "TWO WAYS TO FIX:",
            "  1. Permanent: paste the rows below into inputs/ca_overrides.csv",
            "  2. One-shot : save them to a CSV and re-run with",
            "                python -m oefof_pl --apply-suggestions <file>.csv",
        ]
        if ca_suggestions_path is not None:
            msg_lines.append("")
            msg_lines.append(
                f"Auto-generated file ready to use: {ca_suggestions_path}"
            )
            msg_lines.append(
                f"  Quick re-run: python -m oefof_pl --apply-suggestions \"{ca_suggestions_path}\""
            )
        for line in msg_lines:
            ws.cell(r, 1, line).font = FONT_BODY
            r += 1
        r += 1
        sugg_headers = ("ISIN", "UNITS (signed Diff)", "Suggested REASON (review!)")
        for i, h in enumerate(sugg_headers, start=1):
            cc = ws.cell(r, i, h)
            cc.font = FONT_HEADER; cc.fill = FILL_HEADER; cc.alignment = ALIGN_CENTER
        r += 1
        for f in val01.failures:
            isin = str(f.get("ISIN", ""))
            diff = f.get("Diff", 0.0)
            sname = str(f.get("Stock Name", "")).strip()
            ws.cell(r, 1, isin).font = FONT_BODY
            uc = ws.cell(r, 2, float(diff)); uc.font = FONT_BODY
            uc.number_format = fmt_for("float")
            ws.cell(r, 3, f"{sname} - TODO confirm with ops").font = FONT_BODY
            r += 1
        r += 2

    for i, w in enumerate((10, 28, 8, 60, 10), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_methodology(wb: Workbook, meta: PeriodMeta, *,
                       applied_ca_overrides: tuple = (),
                       ca_override_sources: tuple = (),
                       analyst_sname_overrides: dict | None = None) -> None:
    ws = wb.create_sheet("Methodology")
    _setup_sheet(ws, tab_color=COLOR_TAB_AUDIT)
    _write_title_block(ws, meta,
                       title="Methodology & Audit Trail",
                       subtitle="Read this once. Then read the validation report.")

    text = [
        ("ORD positions", "bold"),
        ("  Total P&L (EUR) = (End MV − Start MV) + Sell Proceeds − Buy Cost + Income", "body"),
        ("  Cost basis uses Weighted Average Cost (WAC). Sells reduce held units at WAC; "
         "the WAC itself is not changed by sells.", "body"),
        ("  Realised + Unrealised + Income does NOT equal Total. Unrealised is reported "
         "point-in-time (End MV − End WAC Cost), not the period change in unrealised. "
         "The cash-flow Total above is the authoritative period P&L.", "body"),
        ("  Cross-zero sells (long → short in one trade) book the entire realised P&L at "
         "the prior WAC and open the residual short at a negative cost basis. VAL-06 "
         "will WARN; review those rows manually if any appear.", "body"),
        ("", "body"),
        ("SWAP / FTSWAP positions", "bold"),
        ("  Realised  = 0", "body"),
        ("  Unrealised = 0", "body"),
        ("  Total P&L (EUR) = Income (EUR)", "body"),
        ("  Cost Basis (EUR) = |Notional × Units| converted at end FX", "body"),
        ("", "body"),
        ("FX policy", "bold"),
        ("  All monetary components (start MV, end MV, buy, sell, income) are converted to EUR "
         "at the END-period FX rate. This eliminates phantom corporate-action gains that arose "
         "in the legacy system from per-trade USD conversion (BUG-4).", "body"),
        ("", "body"),
        ("Closed positions", "bold"),
        ("  When a position is fully exited inside the period the end-period FX may be missing for "
         "its currency. We then fall back to the start-period FX rate so the realised P&L is still "
         "computed in EUR. (Regression covered by tests/test_pl.py.)", "body"),
        ("", "body"),
        ("Corporate-action trades (BCODE = ZZZZ)", "bold"),
        ("  HiPort marks SCODE migrations and bonus-share events with BCODE='ZZZZ' and price=0. "
         "These are NOT economic transactions; we drop them before WAC. Failing to drop them "
         "would dilute per-share cost to zero and book phantom losses on the matched sell.", "body"),
        ("", "body"),
        ("Manual CA overrides (BCODE = CA_ADJ)", "bold"),
        ("  When HiPort's snapshot reflects a corporate action (bonus, split, consolidation, "
         "in-kind delivery) that did NOT arrive in the Bottler trade feed, the unit drift is "
         "patched via inputs/ca_overrides.csv. Each row is injected as a synthetic T='P' trade "
         "at price 0 dated end-period. WAC handles bonus (units up, cost unchanged) and "
         "consolidation (units down, cost unchanged) correctly. Audit trail = the CSV itself; "
         "synthetic rows carry negative TRADE_IDs and BCODE='CA_ADJ' in Trade Detail.", "body"),
        ("", "body"),
        ("Income column = Income & Financing", "bold"),
        ("  HiPort tags every cash-flow leg with T='I' regardless of whether it is a real "
         "dividend or a CFD/swap financing/borrow charge. Negative values in the Income column "
         "are therefore NOT 'negative dividends' \u2014 they are financing legs on CFD/swap "
         "positions (Goldman/HSBC swaps, MSCI ETF CFDs, etc.). The column is labelled "
         "'Income & Financing (EUR)' to make this explicit. Net of the two flows is the cash "
         "P&L impact regardless of category.", "body"),
        ("", "body"),
        ("Percentage handling", "bold"),
        ("  All percent columns are stored as ratios (0.44 = 44%) and Excel's '0.00%' format does "
         "the display conversion. The validate stage rejects anything outside ±5 (the legacy "
         "×100 bug, BUG-1).", "body"),
        ("", "body"),
        ("Inputs", "bold"),
        (f"  HiPort snapshot: {meta.snapshot_path}", "body"),
        (f"  Bottler trades : {meta.bottler_path}", "body"),
        (f"  Generated      : {meta.run_timestamp:%Y-%m-%d %H:%M:%S}", "body"),
        *([
            ("", "body"),
            ("CA Override Sources", "bold"),
            *[(f"  {p}", "body") for p in ca_override_sources],
            (f"  Applied overrides: {len(applied_ca_overrides)} row(s)", "body"),
        ] if ca_override_sources else []),
        *([
            ("", "body"),
            ("Analyst Name Overrides", "bold"),
            *[(f"  {k} → {v}", "body") for k, v in (analyst_sname_overrides or {}).items()],
        ] if analyst_sname_overrides else []),
    ]
    r = 4
    for line, kind in text:
        cell = ws.cell(r, 1, line)
        if kind == "bold":
            cell.font = FONT_SUBTITLE
        else:
            cell.font = FONT_BODY
        cell.alignment = ALIGN_LEFT
        r += 1
    ws.column_dimensions["A"].width = 130
