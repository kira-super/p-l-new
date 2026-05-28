"""Outlook (COM) email sender.

The module is split into two layers:

**Pure helpers** (no COM)
    ``normalize_recipients``, ``build_summary_html_table``,
    ``build_email_html_body`` — all unit-tested without a running Outlook.

**COM-gated send**
    ``send_via_outlook`` — wraps Outlook with an injected ``dispatcher``
    so tests can swap a fake. ``win32com`` is imported only inside the
    default dispatcher function. This module is on the Phase-13 AST lint
    allow-list.

Frozen dataclasses (``EmailMessage``, ``EmailSendResult``) ensure callers
build messages declaratively and that the result of a send is auditable
(account used, attachment count, success flag, error string if any).
"""

from __future__ import annotations

import html
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable, Sequence


# ─── Frozen records ──────────────────────────────────────────────────────────

@dataclass(frozen=True)
class EmailMessage:
    """Declarative email payload (no COM types — pure stdlib)."""
    to: str                            # already normalised (semicolon-separated)
    subject: str
    body_text: str                     # plain-text fallback
    body_html: str | None = None
    cc: str = ""
    attachments: tuple[Path, ...] = ()
    from_smtp: str = ""                # optional Outlook sender account


@dataclass(frozen=True)
class EmailSendResult:
    """Outcome of a `send_via_outlook` call."""
    sent: bool
    to: str
    cc: str
    subject: str
    attachment_count: int
    account_used: str = ""
    error: str = ""


# ─── Pure helpers ────────────────────────────────────────────────────────────

def normalize_recipients(raw_value: object) -> str:
    """Normalise comma/semicolon-separated recipients to a single ``;``-joined string.

    Empty / None returns ``""``. Whitespace is stripped from each token.
    """
    if not raw_value:
        return ""
    tokens = [part.strip() for part in re.split(r"[;,]", str(raw_value)) if part.strip()]
    return ";".join(tokens)


def _format_summary_value(cell) -> str:
    """Display formatter for one Summary-sheet cell value."""
    value = cell.value
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number_format = str(getattr(cell, "number_format", "") or "")
        neg = value < 0
        mag = abs(value)
        if "%" in number_format:
            body = f"{mag * 100:.2f}%"
        elif "€" in number_format or "EUR" in number_format:
            body = f"€{mag:,.0f}"
        elif number_format == "#,##0" or "0;" in number_format and "." not in number_format:
            body = f"{mag:,.0f}"
        else:
            body = f"{mag:,.2f}"
        return f"({body})" if neg else body
    return str(value)


def _scan_table(ws, header_row: int, max_cols: int = 12):
    """Read a contiguous table starting at `header_row` (header + body until
    a blank first-column row). Returns ``(headers, rows)`` where each row is
    a list of cell objects. Stops before the next section header (a row
    whose only populated cell is column 1 with no neighbours)."""
    headers = []
    for c in range(1, max_cols + 1):
        v = ws.cell(header_row, c).value
        if v in (None, ""):
            break
        headers.append(str(v))
    rows = []
    r = header_row + 1
    while True:
        first = ws.cell(r, 1).value
        if first in (None, ""):
            break
        # Detect a section-header row (only col-1 populated).
        next_v = ws.cell(r, 2).value
        if next_v in (None, "") and ws.cell(r, 3).value in (None, ""):
            break
        rows.append([ws.cell(r, c) for c in range(1, len(headers) + 1)])
        r += 1
    return headers, rows


def _find_section_row(ws, label: str, max_row: int = 80) -> int | None:
    """Return the 1-based row index where col A starts with `label` (case
    & whitespace insensitive), else ``None``."""
    target = label.strip().lower()
    for r in range(1, max_row + 1):
        v = ws.cell(r, 1).value
        if v is None:
            continue
        if str(v).strip().lower().startswith(target):
            return r
    return None


def _render_html_table(headers, rows, *, accent_color: str = "#2e6168") -> str:
    """Generic table renderer used for both Instrument Breakdown and
    By Analyst sections in the email body."""
    parts = [
        '<table style="border-collapse:collapse;font-family:Calibri,Arial,sans-serif;'
        'font-size:10.5pt;color:#1f2937;margin:6px 0 16px 0;">',
        f'<thead><tr style="background-color:{accent_color};color:#ffffff;">',
    ]
    for header_text in headers:
        parts.append(
            '<th style="border:1px solid #cbd5d8;padding:6px 10px;'
            'text-align:center;font-weight:700;white-space:nowrap;">'
            f"{html.escape(header_text)}</th>"
        )
    parts.append("</tr></thead><tbody>")
    for r_idx, row in enumerate(rows):
        first_text = str(row[0].value or "").strip().upper() if row else ""
        is_total = first_text == "TOTAL"
        bg = "#dde7e9" if is_total else ("#f5f8f9" if r_idx % 2 == 1 else "#ffffff")
        weight = "700" if is_total else "400"
        border_top = "border-top:2px solid #2e6168;" if is_total else ""
        parts.append(f'<tr style="background-color:{bg};{border_top}">')
        for idx, cell in enumerate(row):
            text = _format_summary_value(cell)
            value = cell.value
            align = "left" if idx == 0 else "right"
            color = "#1f2937"
            if (
                idx != 0
                and isinstance(value, (int, float))
                and not isinstance(value, bool)
                and value < 0
            ):
                color = "#b83a3a"
            parts.append(
                '<td style="border:1px solid #cbd5d8;padding:5px 10px;'
                f'text-align:{align};font-weight:{weight};color:{color};'
                'white-space:nowrap;">'
                f"{html.escape(text)}</td>"
            )
        parts.append("</tr>")
    parts.append("</tbody></table>")
    return "".join(parts)


def _render_kpi_grid(ws) -> str:
    """KPI cards (rows 4–11, cols 1–6) → 4-row, 3-column HTML grid."""
    card_rows = (4, 6, 8, 10)
    card_cols = (1, 3, 5)
    parts = [
        '<table style="border-collapse:separate;border-spacing:8px;'
        'font-family:Calibri,Arial,sans-serif;margin:0 0 14px 0;">'
    ]
    for r in card_rows:
        parts.append("<tr>")
        for c in card_cols:
            label = str(ws.cell(r, c).value or "").strip()
            val_cell = ws.cell(r + 1, c)
            text = _format_summary_value(val_cell)
            value = val_cell.value
            color = "#1f2937"
            if (isinstance(value, (int, float)) and not isinstance(value, bool)
                    and value < 0):
                color = "#b83a3a"
            parts.append(
                '<td style="background-color:#f5f8f9;border:1px solid #cbd5d8;'
                'padding:10px 14px;min-width:170px;">'
                f'<div style="font-size:9pt;color:#5c6b6f;letter-spacing:.4px;'
                f'text-transform:uppercase;">{html.escape(label)}</div>'
                f'<div style="font-size:14pt;font-weight:700;color:{color};'
                f'margin-top:4px;">{html.escape(text)}</div>'
                "</td>"
            )
        parts.append("</tr>")
    parts.append("</table>")
    return "".join(parts)


def build_summary_html_table(summary_workbook_path: Path | str) -> str | None:
    """Build a richly-formatted HTML rendering of the report Summary sheet.

    Extracts three sections:
      1. KPI cards (Total Positions, MV start/end, P&L breakdown, …).
      2. Instrument Breakdown table.
      3. By Analyst — Exposure & P&L table.

    Returns ``None`` if the workbook does not exist or has no Summary sheet.
    """
    path = Path(summary_workbook_path)
    if not path.exists():
        return None

    try:
        from openpyxl import load_workbook
    except ImportError:   # pragma: no cover - dependency present in package
        return None

    wb = load_workbook(path, data_only=True)
    try:
        if "Summary" not in wb.sheetnames:
            return None
        ws = wb["Summary"]

        kpi_html = _render_kpi_grid(ws)

        sections_html: list[str] = []
        for label, accent in (
            ("Instrument Breakdown", "#2e6168"),
            ("By Analyst",           "#3a5a7a"),
        ):
            hdr_row = _find_section_row(ws, label)
            if hdr_row is None:
                continue
            # Section title sits one row above the table header.
            headers, rows = _scan_table(ws, hdr_row + 1)
            if not headers or not rows:
                continue
            title = str(ws.cell(hdr_row, 1).value or label).strip()
            sections_html.append(
                f'<h3 style="font-family:Calibri,Arial,sans-serif;color:#2e6168;'
                f'margin:14px 0 4px 0;font-size:13pt;">{html.escape(title)}</h3>'
                + _render_html_table(headers, rows, accent_color=accent)
            )
    finally:
        wb.close()

    return kpi_html + "".join(sections_html)


def build_email_html_body(
    plain_body: str,
    summary_table_html: str | None,
    *,
    footer_lines: Sequence[str] = (),
) -> str:
    """Wrap `plain_body` in HTML, optionally with summary table + footer."""
    paragraphs = [
        part.strip()
        for part in str(plain_body).replace("\r\n", "\n").split("\n\n")
        if part.strip()
    ]
    parts = [
        '<html><body style="font-family:Calibri,Arial,sans-serif;'
        'font-size:11pt;color:#1f2937;background-color:#ffffff;">'
    ]
    for paragraph in paragraphs:
        parts.append(
            '<p style="margin:0 0 14px 0;line-height:1.45;">'
            f"{html.escape(paragraph).replace(chr(10), '<br>')}"
            "</p>"
        )
    if summary_table_html:
        parts.append(
            '<h2 style="margin:14px 0 6px 0;font-size:15pt;color:#2e6168;'
            'border-bottom:2px solid #2e6168;padding-bottom:4px;">'
            'Portfolio Summary</h2>'
        )
        parts.append(summary_table_html)
        parts.append('<div style="height:14px;"></div>')

    if footer_lines:
        parts.append(
            '<div style="margin-top:14px;padding-top:8px;border-top:1px solid #cbd5d8;'
            'font-size:9pt;color:#5c6b6f;line-height:1.5;">'
            + "<br>".join(html.escape(line) for line in footer_lines)
            + "</div>"
        )
    parts.append("</body></html>")
    return "".join(parts)


# ─── COM-gated send ──────────────────────────────────────────────────────────

def _default_dispatcher(prog_id: str):   # pragma: no cover - COM-only
    """Real Outlook dispatcher (function-local pywin32 import)."""
    import win32com.client as win32
    return win32.Dispatch(prog_id)


def _find_outlook_account(outlook, smtp: str):
    """Return the Outlook Account whose SmtpAddress matches `smtp`, else None."""
    if not smtp:
        return None
    target = str(smtp).strip().lower()
    if not target:
        return None
    try:
        accounts = outlook.Session.Accounts
    except Exception:
        return None
    for account in accounts:
        try:
            if str(account.SmtpAddress).strip().lower() == target:
                return account
        except Exception:
            continue
    return None


def send_via_outlook(
    message: EmailMessage,
    *,
    dispatcher: Callable[[str], object] = _default_dispatcher,
) -> EmailSendResult:
    """Send `message` through Outlook (COM). Returns an audit record.

    The function never raises — failures are reported via the ``error``
    field of :class:`EmailSendResult`. This matches the legacy contract:
    a missing pywin32 install or unreachable Outlook should be a degraded
    pipeline step, not a crash.

    `dispatcher` is injected so tests can supply a fake Outlook object.
    """
    if not message.to:
        return EmailSendResult(
            sent=False, to="", cc=message.cc, subject=message.subject,
            attachment_count=0, error="no recipients",
        )
    existing = tuple(p for p in message.attachments if Path(p).exists())
    if not existing:
        return EmailSendResult(
            sent=False, to=message.to, cc=message.cc, subject=message.subject,
            attachment_count=0, error="no attachments found on disk",
        )

    try:
        outlook = dispatcher("Outlook.Application")
    except ImportError:
        return EmailSendResult(
            sent=False, to=message.to, cc=message.cc, subject=message.subject,
            attachment_count=len(existing), error="pywin32 not installed",
        )
    except Exception as exc:
        return EmailSendResult(
            sent=False, to=message.to, cc=message.cc, subject=message.subject,
            attachment_count=len(existing), error=f"dispatch failed: {exc}",
        )

    try:
        mail = outlook.CreateItem(0)   # 0 = olMailItem
        mail.To = message.to
        if message.cc:
            mail.CC = message.cc
        mail.Subject = message.subject
        mail.Body = message.body_text
        if message.body_html:
            mail.HTMLBody = message.body_html

        account_used = ""
        send_account = _find_outlook_account(outlook, message.from_smtp)
        if send_account is not None:
            mail.SendUsingAccount = send_account
            try:
                account_used = str(send_account.SmtpAddress)
            except Exception:
                account_used = message.from_smtp

        for attachment in existing:
            mail.Attachments.Add(str(attachment))
        mail.Send()
    except Exception as exc:
        return EmailSendResult(
            sent=False, to=message.to, cc=message.cc, subject=message.subject,
            attachment_count=len(existing), account_used="",
            error=f"send failed: {exc}",
        )

    return EmailSendResult(
        sent=True, to=message.to, cc=message.cc, subject=message.subject,
        attachment_count=len(existing), account_used=account_used,
    )
