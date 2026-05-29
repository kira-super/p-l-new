from __future__ import annotations

from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ...compute.pairs import PairMetrics
from ..styles import (
    ALIGN_CENTER,
    ALIGN_LEFT,
    FILL_BAND,
    FILL_HEADER,
    FILL_HEADER_ACCENT,
    FILL_TOTAL,
    FONT_BODY,
    FONT_HEADER,
    FONT_NEG,
    FONT_NEG_BOLD,
    FONT_SECTION,
    FONT_TOTAL,
    align_for,
    fmt_for,
)


FILL_RATIO_HEADER = PatternFill("solid", fgColor="7A8B8D")


def _mv_or_notional_end(instrument: str, cost_basis_eur: float, market_value_end_eur: float) -> float:
    if str(instrument).strip().upper() == "FTSWAP" and abs(float(cost_basis_eur)) > 1e-9:
        return float(cost_basis_eur)
    return float(market_value_end_eur)


def _write_pair_value_row(
    ws: Worksheet,
    row: int,
    col_start: int,
    values: tuple[tuple[object, str], ...],
    *,
    fill=None,
    total: bool = False,
) -> None:
    for idx, (value, kind) in enumerate(values):
        cell = ws.cell(row, col_start + idx, value)
        if kind == "text":
            cell.font = FONT_TOTAL if total else FONT_BODY
        elif isinstance(value, (int, float)) and value < 0:
            cell.font = FONT_NEG_BOLD if total else FONT_NEG
        else:
            cell.font = FONT_TOTAL if total else FONT_BODY
        cell.alignment = align_for(kind)
        if kind != "text":
            cell.number_format = fmt_for(kind)
        if fill is not None:
            cell.fill = fill


def write_pairs_section(
    ws: Worksheet,
    pairs: list[PairMetrics],
    analyst_code: str,
    start_row: int,
    col_start: int = 2,
) -> int:
    analyst_pairs = [p for p in pairs if p.analyst == str(analyst_code).strip().upper()]
    if not analyst_pairs:
        return start_row

    row = start_row
    span = 11
    for col in range(col_start, col_start + span):
        cell = ws.cell(row, col)
        cell.fill = FILL_HEADER
        cell.font = FONT_SECTION
        cell.alignment = ALIGN_LEFT
    ws.cell(row, col_start, "  Pairs & Arb Trades")
    row += 1

    headers = (
        "Leg", "Stock Name", "Units", "Avg Buy Px (EUR)", "Avg Sell Px (EUR)",
        "Current Price (Local)", "Realised (EUR)", "Unrealised (EUR)",
        "Swap Fin (EUR)", "Total P&L (EUR)", "MV / Notional (EUR)",
    )

    for pair in analyst_pairs:
        ws.merge_cells(start_row=row, start_column=col_start,
                       end_row=row, end_column=col_start + span - 1)
        for col in range(col_start, col_start + span):
            cell = ws.cell(row, col)
            cell.fill = FILL_HEADER_ACCENT
            cell.font = FONT_SECTION
            cell.alignment = ALIGN_LEFT
        ws.cell(row, col_start, f"  {pair.label}")
        row += 1

        for idx, header in enumerate(headers):
            cell = ws.cell(row, col_start + idx, header)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER_ACCENT
            cell.alignment = ALIGN_CENTER
        row += 1

        long_values = (
            ("Long", "text"),
            (pair.long_leg.stock_name or pair.long_leg.isin, "text"),
            (pair.long_leg.units, "float"),
            (pair.long_leg.avg_buy_price_eur, "price"),
            (pair.long_leg.avg_sell_price_eur, "price"),
            (pair.long_leg.last_price_eur, "price"),
            (pair.long_leg.realised_pl_eur, "eur"),
            (pair.long_leg.unrealised_pl_eur, "eur"),
            (pair.long_leg.swap_financing_eur, "eur"),
            (pair.long_leg.total_pl_eur, "eur"),
            (pair.long_leg.market_value_end_eur, "eur"),
        )
        _write_pair_value_row(ws, row, col_start, long_values)
        row += 1

        short_values = (
            ("Short", "text"),
            (pair.short_leg.stock_name or pair.short_leg.isin, "text"),
            (pair.short_leg.units, "float"),
            (pair.short_leg.avg_buy_price_eur, "price"),
            (pair.short_leg.avg_sell_price_eur, "price"),
            (pair.short_leg.last_price_eur, "price"),
            (pair.short_leg.realised_pl_eur, "eur"),
            (pair.short_leg.unrealised_pl_eur, "eur"),
            (pair.short_leg.swap_financing_eur, "eur"),
            (pair.short_leg.total_pl_eur, "eur"),
            (_mv_or_notional_end(
                pair.short_leg.instrument,
                pair.short_leg.cost_basis_eur,
                pair.short_leg.market_value_end_eur,
            ), "eur"),
        )
        _write_pair_value_row(ws, row, col_start, short_values, fill=FILL_BAND)
        row += 1

        short_mv_or_notional = _mv_or_notional_end(
            pair.short_leg.instrument,
            pair.short_leg.cost_basis_eur,
            pair.short_leg.market_value_end_eur,
        )
        combined_values = (
            ("Combined", "text"),
            (None, "text"),
            (None, "text"),
            (None, "text"),
            (None, "text"),
            (None, "text"),
            (pair.long_leg.realised_pl_eur + pair.short_leg.realised_pl_eur, "eur"),
            (pair.long_leg.unrealised_pl_eur + pair.short_leg.unrealised_pl_eur, "eur"),
            (pair.long_leg.swap_financing_eur + pair.short_leg.swap_financing_eur, "eur"),
            (pair.combined_total_pl_eur, "eur"),
            (pair.long_leg.market_value_end_eur + short_mv_or_notional, "eur"),
        )
        _write_pair_value_row(ws, row, col_start, combined_values, fill=FILL_TOTAL, total=True)
        row += 1

        if pair.ratio_start is None and pair.ratio_delta is None:
            monitor_headers = ("Current Ratio (Long MV / Short Notional)",)
            monitor_values = ((pair.ratio_end, "float"),)
        else:
            monitor_headers = (
                "Current Ratio (Long MV / Short Notional)",
                "Opening Ratio (L/S)",
                "Ratio Change",
            )
            monitor_values = (
                (pair.ratio_end, "float"),
                (pair.ratio_start, "float"),
                (pair.ratio_delta, "float"),
            )
        for idx, header in enumerate(monitor_headers):
            cell = ws.cell(row, col_start + idx, header)
            cell.font = FONT_HEADER
            cell.fill = FILL_RATIO_HEADER
            cell.alignment = ALIGN_CENTER
        row += 1
        for idx, (value, kind) in enumerate(monitor_values):
            cell = ws.cell(row, col_start + idx, value)
            if isinstance(value, (int, float)) and value < 0:
                cell.font = FONT_NEG
            else:
                cell.font = FONT_BODY
            cell.alignment = align_for(kind)
            if kind != "text":
                cell.number_format = fmt_for(kind)
            cell.fill = FILL_BAND
        row += 2

    return row