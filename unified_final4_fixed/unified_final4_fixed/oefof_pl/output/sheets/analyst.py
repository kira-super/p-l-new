"""Per-analyst sheet writer."""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook

from ...compute.exposure import ExposureMetrics, PeriodMeta
from ..styles import (
    ALIGN_CENTER,
    COLOR_TAB_ANALYST,
    COLUMN_LAYOUT_ANALYST,
    FILL_SUBTOTAL,
    FILL_TOTAL,
    FONT_BODY_MUTED,
    FONT_TOTAL,
)
from ._helpers import (
    _analyst_display,
    _coerce_value,
    _exposure_for,
    _kpi_card,
    _pct_or_none,
    _safe_sheet_name,
    _setup_sheet,
    _with_total_contribution_columns,
    _write_exposure_block_table,
    _write_layout_data,
    _write_layout_headers,
    _write_layout_total,
    _write_section_header,
    _write_title_block,
    get_column_letter,
)


def _write_analyst_sheet(wb: Workbook, pl_df: pd.DataFrame,
                         analyst_code: str, meta: PeriodMeta) -> None:
    name = _safe_sheet_name(analyst_code)
    ws = wb.create_sheet(name)
    _setup_sheet(ws, tab_color=COLOR_TAB_ANALYST, zoom=100)
    _write_title_block(ws, meta,
                       title=f"Analyst: {_analyst_display(analyst_code)} ({analyst_code})",
                       subtitle=(f"Period: {meta.start_date.date():%d %B %Y} → "
                                 f"{meta.end_date.date():%d %B %Y}  |  Fund Currency: EUR"),
                       col_start=2)

    sub = pl_df[pl_df["Analyst"].astype(str).str.upper() == analyst_code].copy()
    cur = sub[sub["Ending Units"] != 0].reset_index(drop=True)
    exi = sub[sub["Ending Units"] == 0].reset_index(drop=True)

    if "Total P&L (EUR)" in cur.columns:
        cur = cur.sort_values("Total P&L (EUR)", ascending=False, na_position="last").reset_index(drop=True)
    if "Total P&L (EUR)" in exi.columns:
        exi = exi.sort_values("Total P&L (EUR)", ascending=False, na_position="last").reset_index(drop=True)
    cur = _with_total_contribution_columns(cur)
    exi = _with_total_contribution_columns(exi)
    sub = _with_total_contribution_columns(sub)
    layout = COLUMN_LAYOUT_ANALYST
    n_cols = len(layout)

    for i, spec in enumerate(layout, start=1):
        ws.column_dimensions[get_column_letter(i)].width = spec.width

    isin_col_idx = next(
        (i for i, spec in enumerate(layout, start=1) if spec.df_column == "ISIN"),
        None,
    )
    if isin_col_idx is not None:
        ws.column_dimensions[get_column_letter(isin_col_idx)].hidden = True

    e = _exposure_for(sub)
    e_cur = _exposure_for(cur)
    em_ytd = (
        meta.ytd_weighted_exposure.by_analyst.get(analyst_code, ExposureMetrics())
        if meta.ytd_weighted_exposure is not None
        else ExposureMetrics()
    )
    ret_on_gross_ytd = (
        _pct_or_none(e["total"], em_ytd.gross_eur)
        if meta.ytd_weighted_exposure is not None and abs(em_ytd.gross_eur) > 1e-9
        else None
    )
    kpi_row = 4
    _kpi_card(ws, row=kpi_row, col=2, label="Period — Positions",
              value=int(len(cur)), kind="int", accent=False)
    _kpi_card(ws, row=kpi_row, col=4, label="Period — Total P&L (EUR)",
              value=e["total"], kind="eur", colour_signed=True, accent=True)
    _kpi_card(ws, row=kpi_row, col=6, label="Period — Total P&L %",
              value=e["total_pct"], kind="pct", colour_signed=True, accent=True)
    _kpi_card(ws, row=kpi_row, col=8, label="Return on Gross Exposure (YTD)",
              value=ret_on_gross_ytd, kind="pct", colour_signed=True, accent=True)
    _kpi_card(ws, row=kpi_row, col=10, label="Current — Net Exposure (EUR)",
              value=e_cur["net_eur"], kind="eur", accent=False)
    _kpi_card(ws, row=kpi_row, col=12, label="Current — Gross Exposure (EUR)",
              value=e_cur["gross_eur"], kind="eur", accent=False)

    visible_span = n_cols - 1
    section_row = kpi_row + 4
    _write_section_header(ws, section_row, title="Current Holdings",
                          span_cols=visible_span, col_start=2)
    header_row = section_row + 1
    _write_layout_headers(ws, header_row, layout)
    next_row = header_row + 1
    if cur.empty:
        empty = ws.cell(next_row, 2, "(none)")
        empty.font = FONT_BODY_MUTED
        next_row += 1
    else:
        next_row = _write_layout_data(ws, cur, layout, start_row=next_row)
        next_row = _write_layout_total(ws, cur, layout, total_row=next_row,
                                        label="Subtotal — Current",
                                        fill=FILL_SUBTOTAL)
    next_row += 1

    ws.merge_cells(start_row=next_row, start_column=2,
                   end_row=next_row, end_column=n_cols)
    total_cell = ws.cell(next_row, 2, "Total")
    total_cell.font = FONT_TOTAL
    total_cell.fill = FILL_TOTAL
    total_cell.alignment = ALIGN_CENTER
    ws.row_dimensions[next_row].height = 20
    next_row += 1

    _write_section_header(ws, next_row, title="Exited Positions",
                          span_cols=visible_span, accent=True, col_start=2)
    header_row = next_row + 1
    EXITED_SUPPRESS = frozenset({"Ending Units", "Last Price (EUR)"})
    _write_layout_headers(ws, header_row, layout, accent=True,
                          suppress=EXITED_SUPPRESS)
    next_row = header_row + 1
    if exi.empty:
        empty = ws.cell(next_row, 2, "(none)")
        empty.font = FONT_BODY_MUTED
        next_row += 1
    else:
        next_row = _write_layout_data(ws, exi, layout, start_row=next_row,
                                       suppress=EXITED_SUPPRESS)
        next_row = _write_layout_total(ws, exi, layout, total_row=next_row,
                                        label="Subtotal — Exited",
                                        fill=FILL_SUBTOTAL,
                                        suppress=EXITED_SUPPRESS)
    next_row += 2

    _write_layout_headers(ws, next_row, layout, accent=True)
    next_row += 1
    TOTAL_SUPPRESS = frozenset({"CCY", "Instrument", "L/S",
                                 "Ending Units", "Last Price (EUR)",
                                 "Avg Buy Price (EUR)",
                                 "Avg Sell Price (EUR)"})
    next_row = _write_layout_total(ws, sub, layout, total_row=next_row,
                                    label="ANALYST TOTAL", fill=FILL_TOTAL,
                                    suppress=TOTAL_SUPPRESS)

    next_row += 2
    if meta.snapshot_exposure is not None:
        next_row = _write_exposure_block_table(
            ws, meta.snapshot_exposure, [analyst_code],
            meta,
            start_row=next_row, is_ytd=False, col_start=2, show_total=False,
        ) + 1
    if meta.ytd_weighted_exposure is not None:
        next_row = _write_exposure_block_table(
            ws, meta.ytd_weighted_exposure, [analyst_code],
            meta,
            start_row=next_row, is_ytd=True, col_start=2, show_total=False,
        ) + 1
