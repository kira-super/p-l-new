"""ValuationA sheet writer."""
from __future__ import annotations

from datetime import datetime
from typing import Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ...compute.exposure import PeriodMeta
from ..styles import (
    ALIGN_CENTER,
    ALIGN_LEFT,
    ALIGN_RIGHT,
    COLOR_TAB_ANALYST,
    FILL_BAND,
    FILL_HEADER,
    FONT_BODY,
    FONT_BODY_MUTED,
    FONT_HEADER,
    FONT_NEG,
    FONT_NEG_BOLD,
    FONT_TOTAL,
    fmt_for,
)
from ._helpers import (
    BORDER_TOP_THICK,
    _coerce_value,
    _safe_float,
    _setup_sheet,
    _write_title_block,
    get_column_letter,
)

# Map raw HiPort pivot header → friendly display label.
_VALA_HEADER_RENAME: dict[str, str] = {
    "SORT1":      "Region",
    "Initials":   "Analyst",
    "SORT2":      "Country",
    "SNAME":      "Stock Name",
    "FMCID":      "Identifier",
    "BTICKER":    "Bloomberg",
    "Cost":       "Cost",
    "Nominal":    "Nominal",
    "Price":      "Price",
    "Fut Value":  "Future Value",
    "P/F Value":  "P/F Value",
    "Exposure %": "Exposure %",
}

# Width hints for renamed ValuationA columns.
_VALA_COL_WIDTHS: dict[str, int] = {
    "Region":       8,
    "Analyst":      9,
    "Country":     12,
    "Stock Name":  28,
    "Total P&L (EUR)": 15,
    "Total P&L %":     11,
    "Identifier":  14,
    "Bloomberg":   18,
    "Cost":        14,
    "Nominal":     12,
    "Price":       10,
    "Future Value": 14,
    "P/F Value":   14,
    "Exposure %":  11,
    "Last Price (EUR)": 13,
    "Avg Buy Price (EUR)": 14,
    "Avg Sell Price (EUR)": 14,
}


def _write_valuation_a(wb: Workbook, valuation_rows: Sequence[Sequence[object]],
                        meta: PeriodMeta, pl_df: pd.DataFrame | None = None, *,
                        th_val_df: pd.DataFrame | None = None) -> None:
    """Render ValuationA in HiPort's hierarchical layout."""
    _ = th_val_df  # reserved
    ws = wb.create_sheet("ValuationA")
    _setup_sheet(ws, tab_color=COLOR_TAB_ANALYST)
    _write_title_block(ws, meta,
                       title="ValuationA — HiPort Reconciliation",
                       subtitle=f"As of {meta.end_date.date():%d %B %Y}")
    if not valuation_rows or len(valuation_rows) < 2:
        ws.cell(4, 1, "(no rows)").font = FONT_BODY_MUTED
        return

    raw_rows = [list(r) for r in valuation_rows]
    hdr_idx = next(
        (i for i, r in enumerate(raw_rows)
         if r and str(r[0] or "").strip().upper() == "SORT1"),
        0,
    )
    raw_headers = raw_rows[hdr_idx]
    raw_body = raw_rows[hdr_idx + 1:]

    def _idx(name: str) -> int | None:
        target = name.upper()
        for i, h in enumerate(raw_headers):
            if str(h or "").strip().upper() == target:
                return i
        return None

    sort1_i = _idx("SORT1")
    init_i = _idx("INITIALS")
    sort2_i = _idx("SORT2")
    sname_i = _idx("SNAME")
    cost_i = _idx("COST")
    nominal_i = _idx("NOMINAL")
    futv_i = _idx("FUT VALUE")
    pfv_i = _idx("P/F VALUE")
    exp_i = _idx("EXPOSURE %")

    sname_lookup: dict[str, tuple[str, float | None, float | None,
                                  float | None, float | None, float | None]] = {}
    if pl_df is not None and not pl_df.empty:
        for _, r in pl_df.iterrows():
            key = str(r.get("Stock Name", "")).strip().upper()
            if not key:
                continue
            analyst = str(r.get("Analyst", "") or "").strip().upper()
            pl_eur = r.get("Total P&L (EUR)")
            pl_pct = r.get("Total P&L (%)")
            last_px = r.get("Last Price (EUR)")
            buy_px = r.get("Avg Buy Price (EUR)")
            sell_px = r.get("Avg Sell Price (EUR)")
            sname_lookup[key] = (
                analyst,
                float(pl_eur) if pd.notna(pl_eur) else None,
                float(pl_pct) if pd.notna(pl_pct) else None,
                float(last_px) if pd.notna(last_px) else None,
                float(buy_px) if pd.notna(buy_px) else None,
                float(sell_px) if pd.notna(sell_px) else None,
            )

    cur_sort1 = cur_init = cur_sort2 = ""
    Stock = dict
    stocks: list[Stock] = []
    for r in raw_body:
        if not any(v not in (None, "") for v in r):
            continue
        s1 = str(r[sort1_i] or "").strip() if sort1_i is not None and sort1_i < len(r) else ""
        ini = str(r[init_i] or "").strip() if init_i is not None and init_i < len(r) else ""
        s2 = str(r[sort2_i] or "").strip() if sort2_i is not None and sort2_i < len(r) else ""
        sn = str(r[sname_i] or "").strip() if sname_i is not None and sname_i < len(r) else ""
        if s1:
            cur_sort1 = s1
        if ini:
            cur_init = ini
        if s2 and "Total" not in s2:
            cur_sort2 = s2
        if not sn:
            continue
        eff_init = cur_init
        if eff_init in ("", "(blank)"):
            mapped = sname_lookup.get(sn.upper())
            if mapped and mapped[0]:
                eff_init = mapped[0]
        pl_eur, pl_pct = (None, None)
        last_px = buy_px = sell_px = None
        hit = sname_lookup.get(sn.upper())
        if hit:
            pl_eur, pl_pct, last_px, buy_px, sell_px = (
                hit[1], hit[2], hit[3], hit[4], hit[5],
            )
        stocks.append({
            "raw": r, "sort1": cur_sort1, "init": eff_init or "(blank)",
            "country": cur_sort2, "sname": sn,
            "pl_eur": pl_eur, "pl_pct": pl_pct,
            "last_px": last_px, "buy_px": buy_px, "sell_px": sell_px,
        })

    _VALA_DROPPED = {"Identifier", "Cost", "Nominal", "Price", "P/F Value"}
    display_headers: list[str] = []
    raw_keep: list[int] = []
    sname_out = None
    for raw_idx, raw_h in enumerate(raw_headers):
        nice = _VALA_HEADER_RENAME.get(str(raw_h or "").strip(),
                                       str(raw_h or "").strip())
        if nice in _VALA_DROPPED:
            continue
        display_headers.append(nice)
        raw_keep.append(raw_idx)
        if nice == "Stock Name":
            sname_out = len(display_headers) - 1
    if sname_out is not None:
        display_headers[sname_out + 1:sname_out + 1] = [
            "Total P&L (EUR)", "Total P&L %",
        ]
    display_headers.extend([
        "Last Price (EUR)", "Avg Buy Price (EUR)", "Avg Sell Price (EUR)",
    ])
    n_cols = len(display_headers)

    header_row = 4
    for i, k in enumerate(display_headers, start=1):
        cc = ws.cell(header_row, i, k)
        cc.font = FONT_HEADER
        cc.fill = FILL_HEADER
        cc.alignment = ALIGN_CENTER
        ws.column_dimensions[get_column_letter(i)].width = _VALA_COL_WIDTHS.get(
            k, max(10, min(20, len(k) + 4)),
        )
    ws.row_dimensions[header_row].height = 24

    def _stock_key(s):
        unassigned = s["init"] in ("", "(blank)")
        return (1 if unassigned else 0, s["init"], s["country"], s["sname"].upper())
    stocks.sort(key=_stock_key)

    raw_to_out: dict[int, int] = {}
    insert_after = sname_out
    _SYNTH_HEADERS = {
        "Total P&L (EUR)", "Total P&L %",
        "Last Price (EUR)", "Avg Buy Price (EUR)", "Avg Sell Price (EUR)",
    }
    keep_iter = iter(raw_keep)
    for out_idx, h in enumerate(display_headers):
        if h in _SYNTH_HEADERS:
            continue
        raw_to_out[next(keep_iter)] = out_idx
    pl_eur_out = (insert_after + 1) if insert_after is not None else None
    pl_pct_out = (insert_after + 2) if insert_after is not None else None
    last_px_out = n_cols - 3
    buy_px_out = n_cols - 2
    sell_px_out = n_cols - 1

    out_init = raw_to_out.get(init_i) if init_i is not None else None
    out_sort1 = raw_to_out.get(sort1_i) if sort1_i is not None else None
    out_sort2 = raw_to_out.get(sort2_i) if sort2_i is not None else None
    out_sname = raw_to_out.get(sname_i) if sname_i is not None else None
    out_cost = raw_to_out.get(cost_i) if cost_i is not None else None
    out_nominal = raw_to_out.get(nominal_i) if nominal_i is not None else None
    out_futv = raw_to_out.get(futv_i) if futv_i is not None else None
    out_pfv = raw_to_out.get(pfv_i) if pfv_i is not None else None
    out_exp = raw_to_out.get(exp_i) if exp_i is not None else None

    FILL_COUNTRY_TOTAL = PatternFill("solid", fgColor="EAEFF1")
    FILL_ANALYST_TOTAL = PatternFill("solid", fgColor="C8D5DA")

    def _format_cell(cell, value, *, is_pct_eur=False, is_pct=False,
                     is_total=False):
        if value is None or value == "":
            cell.value = ""
            cell.alignment = ALIGN_LEFT
        elif isinstance(value, bool):
            cell.value = "Yes" if value else "No"
            cell.alignment = ALIGN_CENTER
        elif isinstance(value, (int, float)):
            cell.value = float(value)
            if is_pct_eur:
                cell.number_format = fmt_for("eur")
            elif is_pct:
                cell.number_format = fmt_for("pct")
            else:
                cell.number_format = "#,##0.00"
            cell.alignment = ALIGN_RIGHT
            if cell.value < 0 and (is_pct_eur or is_pct):
                cell.font = FONT_NEG_BOLD if is_total else FONT_NEG
                return
        elif isinstance(value, datetime):
            cell.value = value
            cell.number_format = "yyyy-mm-dd"
            cell.alignment = ALIGN_CENTER
        else:
            cell.value = str(value)
            cell.alignment = ALIGN_LEFT
        cell.font = FONT_TOTAL if is_total else FONT_BODY

    def _write_subtotal_row(row_idx: int, label: str, label_col: int,
                             stocks_subset, fill, *, bold_label: bool):
        ws.row_dimensions[row_idx].height = 18
        sums = {
            "cost":    sum(_safe_float(s["raw"][cost_i]) for s in stocks_subset)    if cost_i    is not None else 0.0,
            "nominal": sum(_safe_float(s["raw"][nominal_i]) for s in stocks_subset) if nominal_i is not None else 0.0,
            "futv":    sum(_safe_float(s["raw"][futv_i]) for s in stocks_subset)    if futv_i    is not None else 0.0,
            "pfv":     sum(_safe_float(s["raw"][pfv_i]) for s in stocks_subset)     if pfv_i     is not None else 0.0,
            "exp":     sum(_safe_float(s["raw"][exp_i]) for s in stocks_subset)     if exp_i     is not None else 0.0,
            "pl_eur":  sum(s["pl_eur"] for s in stocks_subset if s["pl_eur"] is not None),
        }
        for c_idx in range(n_cols):
            cell = ws.cell(row_idx, c_idx + 1)
            cell.fill = fill
            cell.border = BORDER_TOP_THICK if bold_label else None
            if c_idx == label_col:
                cell.value = label
                cell.font = FONT_TOTAL if bold_label else FONT_BODY
                cell.alignment = ALIGN_LEFT
            elif c_idx == out_cost and out_cost is not None:
                _format_cell(cell, sums["cost"], is_total=True)
            elif c_idx == out_nominal and out_nominal is not None:
                _format_cell(cell, sums["nominal"], is_total=True)
            elif c_idx == out_futv and out_futv is not None:
                _format_cell(cell, sums["futv"], is_total=True)
            elif c_idx == out_pfv and out_pfv is not None:
                _format_cell(cell, sums["pfv"], is_total=True)
            elif c_idx == out_exp and out_exp is not None:
                _format_cell(cell, sums["exp"], is_pct=True, is_total=True)
            elif c_idx == pl_eur_out and pl_eur_out is not None:
                _format_cell(cell, sums["pl_eur"], is_pct_eur=True, is_total=True)
            else:
                cell.value = ""
                cell.font = FONT_TOTAL if bold_label else FONT_BODY
        return row_idx + 1

    out_row = header_row + 1
    prev_init = None
    prev_country_in_init = None
    cur_init_stocks: list = []
    cur_country_stocks: list = []

    def _flush_country(at_row: int) -> int:
        nonlocal prev_country_in_init
        if cur_country_stocks and prev_country_in_init is not None:
            label = f"{prev_country_in_init} Total"
            at_row = _write_subtotal_row(
                at_row, label, out_sort2 if out_sort2 is not None else 0,
                cur_country_stocks, FILL_COUNTRY_TOTAL, bold_label=False,
            )
        return at_row

    def _flush_init(at_row: int) -> int:
        nonlocal prev_init
        if cur_init_stocks and prev_init is not None:
            label = f"{prev_init} Total"
            at_row = _write_subtotal_row(
                at_row, label, out_init if out_init is not None else 0,
                cur_init_stocks, FILL_ANALYST_TOTAL, bold_label=True,
            )
        return at_row

    for s in stocks:
        if s["init"] != prev_init:
            out_row = _flush_country(out_row)
            out_row = _flush_init(out_row)
            prev_init = s["init"]
            prev_country_in_init = None
            cur_init_stocks = []
            cur_country_stocks = []
            show_init = True
        else:
            show_init = False

        if s["country"] != prev_country_in_init:
            out_row = _flush_country(out_row)
            prev_country_in_init = s["country"]
            cur_country_stocks = []
            show_country = True
        else:
            show_country = False

        ws.row_dimensions[out_row].height = 16
        for c_idx in range(n_cols):
            cell = ws.cell(out_row, c_idx + 1)
            if c_idx == out_init:
                cell.value = s["init"] if show_init else ""
                cell.alignment = ALIGN_LEFT
                cell.font = FONT_TOTAL if show_init else FONT_BODY
            elif c_idx == out_sort1:
                cell.value = s["sort1"] if show_init else ""
                cell.alignment = ALIGN_CENTER
                cell.font = FONT_BODY
            elif c_idx == out_sort2:
                cell.value = s["country"] if show_country else ""
                cell.alignment = ALIGN_LEFT
                cell.font = FONT_BODY
            elif c_idx == pl_eur_out:
                _format_cell(cell, s["pl_eur"], is_pct_eur=True)
            elif c_idx == pl_pct_out:
                _format_cell(cell, s["pl_pct"], is_pct=True)
            elif c_idx == last_px_out:
                _format_cell(cell, s["last_px"])
            elif c_idx == buy_px_out:
                _format_cell(cell, s["buy_px"])
            elif c_idx == sell_px_out:
                _format_cell(cell, s["sell_px"])
            else:
                raw_idx = next((ri for ri, oi in raw_to_out.items() if oi == c_idx), None)
                v = s["raw"][raw_idx] if raw_idx is not None and raw_idx < len(s["raw"]) else None
                if c_idx == out_exp:
                    _format_cell(cell, v, is_pct=True)
                else:
                    _format_cell(cell, v)
        cur_country_stocks.append(s)
        cur_init_stocks.append(s)
        out_row += 1

    out_row = _flush_country(out_row)
    out_row = _flush_init(out_row)

    if stocks:
        FILL_GRAND_TOTAL = PatternFill("solid", fgColor="2E6168")
        ws.row_dimensions[out_row].height = 22
        sums = {
            "cost":    sum(_safe_float(s["raw"][cost_i]) for s in stocks)    if cost_i    is not None else 0.0,
            "nominal": sum(_safe_float(s["raw"][nominal_i]) for s in stocks) if nominal_i is not None else 0.0,
            "futv":    sum(_safe_float(s["raw"][futv_i]) for s in stocks)    if futv_i    is not None else 0.0,
            "pfv":     sum(_safe_float(s["raw"][pfv_i]) for s in stocks)     if pfv_i     is not None else 0.0,
            "exp":     sum(_safe_float(s["raw"][exp_i]) for s in stocks)     if exp_i     is not None else 0.0,
            "pl_eur":  sum(s["pl_eur"] for s in stocks if s["pl_eur"] is not None),
        }
        WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for c_idx in range(n_cols):
            cell = ws.cell(out_row, c_idx + 1)
            cell.fill = FILL_GRAND_TOTAL
            cell.border = BORDER_TOP_THICK
            cell.font = WHITE_BOLD
            if c_idx == 0:
                cell.value = "Grand Total"
                cell.alignment = ALIGN_LEFT
            elif c_idx == out_cost and out_cost is not None:
                cell.value = sums["cost"]; cell.number_format = "#,##0.00"; cell.alignment = ALIGN_RIGHT
            elif c_idx == out_nominal and out_nominal is not None:
                cell.value = sums["nominal"]; cell.number_format = "#,##0.00"; cell.alignment = ALIGN_RIGHT
            elif c_idx == out_futv and out_futv is not None:
                cell.value = sums["futv"]; cell.number_format = "#,##0.00"; cell.alignment = ALIGN_RIGHT
            elif c_idx == out_pfv and out_pfv is not None:
                cell.value = sums["pfv"]; cell.number_format = "#,##0.00"; cell.alignment = ALIGN_RIGHT
            elif c_idx == out_exp and out_exp is not None:
                cell.value = sums["exp"]; cell.number_format = fmt_for("pct"); cell.alignment = ALIGN_RIGHT
            elif c_idx == pl_eur_out and pl_eur_out is not None:
                cell.value = sums["pl_eur"]; cell.number_format = fmt_for("eur"); cell.alignment = ALIGN_RIGHT
            else:
                cell.value = ""
        out_row += 1

    ws.freeze_panes = ws.cell(header_row + 1, 1)
